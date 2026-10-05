"""Build the Saffron client delivery sheet (D23–D24 deliverable) from the two Clay exports.

Inputs:  data/saffron_accounts_qualified.csv  (accounts table export)
         data/saffron_people_export.csv       (Saffron People table export)
Output:  data/Saffron - Client Delivery Sheet.xlsx
Tabs:    Summary · Contacts (sendable) · Accounts (all 50) · Follow-ups
"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ACC = "data/saffron_accounts_qualified.csv"
PPL = "data/saffron_people_export.csv"
OUT = "data/Saffron - Client Delivery Sheet.xlsx"

# Waterfall steps in run order: (find column, provider name). Validation columns sit between them.
WATERFALL = [
    ("Find Work Email", "Findymail"), ("Find Work Email (2)", "Hunter"), ("Find Work Email (3)", "Prospeo"),
    ("Find work email", "Kitt"), ("Find Work Email (4)", "Datagma"), ("Find work email (2)", "Wiza"),
    ("Find Work Email (5)", "Icypeas"), ("Find Work Email (6)", "Enrow"), ("Find work email (3)", "Dropcontact"),
    ("Find work email (4)", "LeadMagic"), ("Find Work Email (7)", "SMARTe"),
]
NO_PEOPLE = {"Cognition", "Polymarket", "RoboMQ"}  # Surfe returned nobody matching the title/seniority filters

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="0E6F69")
OK_FILL = PatternFill("solid", fgColor="E2F1EF")
NO_FILL = PatternFill("solid", fgColor="FBF0E3")


def read(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def provider_for(row):
    """The waterfall provider whose find step returned an address that reached the final column."""
    email = row["Work Email"].strip().lower()
    for col, name in WATERFALL:
        if email and email in row.get(col, "").lower():
            return name
    return ""


def write_sheet(ws, header, rows, widths, fill_col=None):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    if fill_col is not None:
        for row in ws.iter_rows(min_row=2):
            v = row[fill_col].value or ""
            fill = OK_FILL if v in ("Outreach", "Yes") else NO_FILL if v else None
            if fill:
                for c in row:
                    c.fill = fill


def main():
    acc, ppl = read(ACC), read(PPL)
    outreach = [a["Company Name"] for a in acc if a["Outreach Eligibility"] == "Outreach"]
    in_people = {p["Company Name"] for p in ppl}
    over_limit = [c for c in outreach if c not in in_people and c not in NO_PEOPLE]

    contacts, followups = [], []
    for p in ppl:
        persona = p["Job Seniority"].strip()
        email = p["Work Email"].strip()
        base = [p["Company Name"], p["Full Name"], p["First Name"], p["Job Title"], persona or "(not a buyer)",
                p["Linked In Url"], email]
        if email:
            contacts.append(base + [provider_for(p), "Valid", p["Qualification Reason"], p["Tools"],
                                    p["Employee Count"], p["Jobcount"], p["Clean Domain"]])
        elif persona:
            found_unvalidated = any(p.get(col, "").strip() and "@" in p.get(col, "") for col, _ in WATERFALL)
            why = ("Address found but failed validation; excluded" if found_unvalidated
                   else "No email from any of 11 providers")
            followups.append(["Reach via LinkedIn", p["Company Name"], p["Full Name"], p["Job Title"],
                              p["Linked In Url"], why])
        else:
            followups.append(["Wrong-fit title (skipped, no credits spent)", p["Company Name"], p["Full Name"],
                              p["Job Title"], p["Linked In Url"], "Sales/customer-facing role returned by the people search"])
    for c in over_limit:
        followups.append(["Next batch (free-plan 50-row limit)", c, "", "", "",
                          "Outreach company; people were found but sit beyond row 50 of the free plan"])
    for c in outreach:
        if c in NO_PEOPLE:
            followups.append(["Find people manually", c, "", "", "",
                              "People search returned nobody matching the buyer titles"])
    followups.append(["Manual review", "Pinecone", "", "", "",
                      "Disqualified only for 0 recruiting team; likely a data gap (9 SWE openings). Check LinkedIn."])
    followups.append(["Manual review", "Avoma", "", "", "",
                      "Mentions AI coding tools but 66 employees and no TA team; kept as Don't outreach."])

    accounts = [[a["Company Name"], a["Clean Domain"], a["Company LinkedIn URL"], a["Employee Count"], a["Industry"],
                 a["Country"], a["Founded"], a["Jobcount"] or "0", a["Peoplecount"] or "0",
                 a["Ai Tools Mentioned"], a["Tools"], a["Outreach Eligibility"],
                 a.get("Qualification Reason") or a.get("Staffing Summary", "")] for a in acc]

    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    summary = [
        ("Saffron outbound pipeline: client delivery sheet", ""),
        ("Built for", "Saffron (YC Spring 2026), AI-native technical assessments"),
        ("Target segment", "US tech companies, 50–1,000 employees, hiring software engineers, with an in-house recruiting team"),
        ("Accounts screened", len(acc)),
        ("Qualified (Outreach)", len(outreach)),
        ("Contacts found (first 50 rows)", len(ppl)),
        ("Verified work emails", len(contacts)),
        ("Follow-up items", len(followups)),
        ("Clay credits used (whole pipeline)", 118.6),
        ("Qualification rule", "Don't outreach if: no recruiting team, <50 or >1,000 employees, non-US, or no SWE openings. "
                               "Outreach if: 5+ SWE openings, or 1–4 openings and job ads name AI coding tools."),
        ("Email policy", "Only addresses that passed the waterfall's validation step are listed."),
    ]
    for k, v in summary:
        s.append([k, v])
    s["A1"].font = Font(bold=True, size=14)
    s.column_dimensions["A"].width = 36
    s.column_dimensions["B"].width = 110

    write_sheet(wb.create_sheet("Contacts"),
                ["Company", "Full Name", "First Name", "Job Title", "Persona", "LinkedIn", "Work Email",
                 "Email found by", "Email status", "Why this account", "AI tools in job ads", "Employees",
                 "SWE openings", "Domain"],
                contacts, [16, 22, 12, 34, 15, 40, 32, 13, 11, 48, 30, 10, 12, 18])
    write_sheet(wb.create_sheet("Accounts"),
                ["Company", "Domain", "Company LinkedIn", "Employees", "Industry", "Country", "Founded",
                 "SWE openings", "Recruiting team", "AI tools mentioned", "Tools", "Eligibility", "Reason"],
                accounts, [22, 20, 40, 10, 26, 8, 9, 12, 14, 12, 30, 14, 60], fill_col=11)
    write_sheet(wb.create_sheet("Follow-ups"),
                ["Action", "Company", "Person", "Title", "LinkedIn", "Detail"],
                followups, [34, 16, 22, 34, 40, 70])
    wb.save(OUT)
    print(f"accounts {len(acc)} | outreach {len(outreach)} | people {len(ppl)} | emails {len(contacts)} | "
          f"follow-ups {len(followups)} | over-limit companies {len(over_limit)}")
    print("provider split:", {n: sum(1 for c in contacts if c[7] == n) for _, n in WATERFALL if any(c[7] == n for c in contacts)})
    print("wrote", OUT)


if __name__ == "__main__":
    main()
