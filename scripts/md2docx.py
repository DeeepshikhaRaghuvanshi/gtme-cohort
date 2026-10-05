"""Minimal Markdown -> .docx converter (no pandoc on this box).
Handles: # headings, paragraphs, - / * bullets (2-level), 1. numbered lists,
| tables |, > quotes, ``` code blocks, **bold**, *italic*, `code`, [text](url).
Usage: python3 scripts/md2docx.py in.md out.docx
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<!\*)\*[^*\s][^*]*\*(?!\*))")


def add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif part.startswith("[") and "](" in part:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            r = par.add_run(f"{label} ({url})")
            r.font.color.rgb = RGBColor(0x0E, 0x6F, 0x69)
        elif part.startswith("*") and part.endswith("*"):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def convert(src, dst):
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or s in ("---", "-----"):
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            level = min(len(m.group(1)), 4)
            h = doc.add_heading(level=0 if level == 1 else level - 1)
            add_inline(h, m.group(2))
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        add_inline(cell.paragraphs[0], row[ci] if ci < len(row) else "")
                        if ri == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
            continue
        m = re.match(r"^(\s*)[-*]\s+(.*)", line)
        if m:
            style = "List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
            add_inline(doc.add_paragraph(style=style), m.group(2))
            i += 1
            continue
        m = re.match(r"^(\s*)\d+[.)]\s+(.*)", line)
        if m:
            style = "List Number 2" if len(m.group(1)) >= 2 else "List Number"
            add_inline(doc.add_paragraph(style=style), m.group(2))
            i += 1
            continue
        if s.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, s.lstrip("> "))
            i += 1
            continue
        add_inline(doc.add_paragraph(), s)
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
